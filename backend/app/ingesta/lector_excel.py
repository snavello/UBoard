"""Excel (.xlsx/.xlsm) -> una tabla DuckDB VARCHAR por hoja con datos.

openpyxl en modo solo lectura entrega las celdas ya tipadas; se pasan a texto
con un formato canonico (fechas ISO, enteros sin ".0") y despues el tipado
comun decide, igual que con un CSV. Asi Excel y CSV siguen un unico camino.
"""
import io
from datetime import date, datetime, time
from pathlib import Path

import duckdb
import pyarrow
from openpyxl import load_workbook

from app.ingesta.encabezado import detectar_fila_encabezado, nombres_de_columnas
from app.ingesta.lector_csv import TablaCruda
from app.nucleo.errores import ErrorApp


def _celda_a_texto(valor) -> str | None:
    if valor is None:
        return None
    if isinstance(valor, bool):
        return "true" if valor else "false"
    if isinstance(valor, datetime):
        if valor.time() == time(0, 0):
            return valor.date().isoformat()
        return valor.isoformat(sep=" ")
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    texto = str(valor).strip()
    return texto or None


def cargar_excel(
    conexion: duckdb.DuckDBPyConnection, datos: bytes, prefijo_relacion: str, directorio_temporal: Path
) -> list[tuple[str, TablaCruda]]:
    """Devuelve [(nombre de hoja, tabla)] solo para las hojas con datos."""
    try:
        libro = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    except Exception as error:  # noqa: BLE001  openpyxl tira de todo (zip roto, xml invalido)
        raise ErrorApp("E-ING-01", f"{type(error).__name__}: {error}") from error
    resultado = []
    for numero, hoja in enumerate(libro.worksheets, start=1):
        if hasattr(hoja, "reset_dimensions"):
            hoja.reset_dimensions()  # dimensiones declaradas a veces mienten
        filas = [[_celda_a_texto(celda) for celda in fila] for fila in hoja.iter_rows(values_only=True)]
        indice, tiene_encabezado = detectar_fila_encabezado(filas)
        if indice is None:
            continue
        cuerpo = filas[indice:]
        ancho = max(len(fila) for fila in cuerpo)
        cuerpo = [fila + [None] * (ancho - len(fila)) for fila in cuerpo]
        origen: list[str | None] = cuerpo[0] if tiene_encabezado else [None] * ancho
        datos_filas = cuerpo[1:] if tiene_encabezado else cuerpo
        datos_filas = [fila for fila in datos_filas if any(celda is not None for celda in fila)]
        if not datos_filas:
            continue
        columnas = nombres_de_columnas(origen)
        tabla_arrow = pyarrow.table(
            {columna: pyarrow.array([fila[posicion] for fila in datos_filas], type=pyarrow.string()) for posicion, columna in enumerate(columnas)}
        )
        nombre_relacion = f"{prefijo_relacion}_h{numero}"
        conexion.register(f"{nombre_relacion}_arrow", tabla_arrow)
        conexion.execute(f"CREATE OR REPLACE TABLE {nombre_relacion} AS SELECT * FROM {nombre_relacion}_arrow")
        conexion.unregister(f"{nombre_relacion}_arrow")
        resultado.append(
            (
                hoja.title,
                TablaCruda(
                    nombre_relacion=nombre_relacion,
                    columnas=columnas,
                    columnas_origen=origen,
                    opciones={"hoja": hoja.title, "filas_saltadas": indice, "tiene_encabezado": tiene_encabezado},
                ),
            )
        )
    libro.close()
    return resultado
